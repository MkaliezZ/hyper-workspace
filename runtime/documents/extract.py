"""Bounded, read-only text extraction. No model, network or source credentials."""
import io
import json
import re
import sys
import zipfile
from pathlib import PurePosixPath

VERSION = "document-text-v1"
MAX_BYTES = 8 * 1024 * 1024
MAX_TEXT = 180000
MAX_SECTIONS = 2000


class ExtractionError(ValueError):
    pass


def check_package(data, kind):
    """Bound decompression before any OOXML library sees the package."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 2048:
                raise ExtractionError("Office 文件部件过多。")
            names, total = set(), 0
            for member in members:
                name = member.filename
                path = PurePosixPath(name)
                if (name in names or path.is_absolute() or ".." in path.parts
                        or "\\" in name or member.flag_bits & 1
                        or (member.external_attr >> 16) & 0o170000 == 0o120000):
                    raise ExtractionError("Office 文件包包含不支持的路径或加密部件。")
                names.add(name)
                total += member.file_size
                if (member.file_size > 8 * 1024 * 1024 or total > 32 * 1024 * 1024
                        or member.file_size > max(member.compress_size, 1) * 200):
                    raise ExtractionError("Office 文件解压后超过解析限额。")
                raw = archive.read(member)
                # XML relationships can target another extension. Binary printer
                # settings may embed a plist, but are never parsed as XML here.
                is_xml = (name.lower().endswith((".xml", ".rels"))
                          or raw.lstrip(b" \t\r\n\xef\xbb\xbf").startswith(b"<")
                          or raw.startswith((b"\xff\xfe", b"\xfe\xff")))
                if is_xml:
                    if re.search(br"<!\s*(?:DOCTYPE|ENTITY)", raw, re.I) or b"\x00" in raw or raw.startswith((b"\xff\xfe", b"\xfe\xff")):
                        raise ExtractionError("Office XML 包含不支持的声明或编码。")
                    if kind == "docx" and name == "word/document.xml":
                        if re.search(br"<(?:\w+:)?(?:ins|del)(?:\s|>)", raw):
                            raise ExtractionError("请先在 Word 中确认修订后再导入。")
            required = {"docx": "word/document.xml", "xlsx": "xl/workbook.xml",
                        "pptx": "ppt/presentation.xml"}[kind]
            if required not in names or "[Content_Types].xml" not in names:
                raise ExtractionError("Office 文件内容与选定格式不符。")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise ExtractionError("当前不支持包含宏的 Office 文件。")
    except (zipfile.BadZipFile, RuntimeError, KeyError, NotImplementedError) as exc:
        raise ExtractionError("Office 文件包损坏或压缩方式不受支持。") from exc


def extract(data, kind):
    if kind not in ("pdf", "docx", "xlsx", "pptx"):
        raise ExtractionError("不支持的文档格式。")
    if not data or len(data) > MAX_BYTES:
        raise ExtractionError("文档为空或超过 8 MiB。")
    sections, warnings, total = [], [], 0

    def add(locator, text):
        nonlocal total
        text = str(text).replace("\v", "\n").replace("\f", "\n").strip()
        if not text:
            return
        if re.search(r"[\x00-\x08\x0e-\x1f]", text):
            raise ExtractionError("文档文本包含不支持的控制字符。")
        # Match the product's UTF-16 positions and limit, including non-BMP characters.
        total += len(locator.encode("utf-16-le")) // 2 + len(text.encode("utf-16-le")) // 2 + 5
        if total > MAX_TEXT or len(sections) >= MAX_SECTIONS:
            raise ExtractionError("文档提取文本超过 180000 字符，请缩小资料范围。")
        sections.append({"locator": locator, "text": text})

    if kind == "pdf":
        if not data.startswith(b"%PDF-"):
            raise ExtractionError("文件内容不是 PDF。")
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise ExtractionError("当前不支持加密 PDF。")
        if len(reader.pages) > 50:
            raise ExtractionError("PDF 超过 50 页。")
        for index, page in enumerate(reader.pages, 1):
            content = page.get_contents()
            if content is not None and len(content.get_data()) > 8 * 1024 * 1024:
                raise ExtractionError("PDF 页面解压内容超过限额。")
            text = page.extract_text() or ""
            if not text.strip():
                warnings.append(f"第 {index} 页没有可提取文本，未执行 OCR。")
            add(f"第 {index} 页", text)
    else:
        check_package(data, kind)
        if kind == "docx":
            from docx import Document
            from docx.table import Table
            document = Document(io.BytesIO(data))

            def blocks(container, prefix, depth=0):
                if depth > 8:
                    raise ExtractionError("Word 表格嵌套超过限额。")
                for index, block in enumerate(container.iter_inner_content(), 1):
                    location = f"{prefix}/块{index}"
                    if isinstance(block, Table):
                        seen = set()
                        for ri, row in enumerate(block.rows, 1):
                            for ci, cell in enumerate(row.cells, 1):
                                if cell._tc in seen:
                                    continue
                                seen.add(cell._tc)
                                blocks(cell, f"{location}/表格R{ri}C{ci}", depth + 1)
                    else:
                        add(location, block.text)
            blocks(document, "Word正文")
            warnings.append("仅提取正文段落与表格；未包含页眉、页脚、批注及图片文字。")
        elif kind == "xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) > 10:
                    raise ExtractionError("Excel 超过 10 个工作表。")
                for sheet in workbook.worksheets:
                    if sheet.sheet_state != "visible":
                        warnings.append(f"未读取隐藏工作表：{sheet.title}")
                        continue
                    # Read-only dimensions are untrusted; bound actual enumeration too.
                    if sheet.max_row and sheet.max_row > 2000 or sheet.max_column and sheet.max_column > 100:
                        raise ExtractionError("工作表超过 2000 行或 100 列。")
                    sheet.reset_dimensions()
                    count = 0
                    for ri, row in enumerate(sheet.iter_rows(), 1):
                        if ri > 2000 or len(row) > 100:
                            raise ExtractionError("工作表实际范围超过限额。")
                        for cell in row:
                            count += 1
                            if count > 50000:
                                raise ExtractionError("工作表单元格超过限额。")
                            if cell.value is not None:
                                prefix = "公式（未计算）：" if cell.data_type == "f" else ""
                                add(f"{sheet.title}!{cell.coordinate}", prefix + str(cell.value))
            finally:
                workbook.close()
            warnings.append("未计算公式，未读取批注、图表或图片文字。")
        else:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            presentation = Presentation(io.BytesIO(data))
            if len(presentation.slides) > 50:
                raise ExtractionError("PowerPoint 超过 50 页。")

            def shapes(values, prefix, depth=0):
                if depth > 8:
                    raise ExtractionError("幻灯片组合嵌套超过限额。")
                for index, shape in enumerate(values, 1):
                    location = f"{prefix}/形状{index}"
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        shapes(shape.shapes, location, depth + 1)
                    elif shape.has_table:
                        for ri, row in enumerate(shape.table.rows, 1):
                            for ci, cell in enumerate(row.cells, 1):
                                if not cell.is_spanned:
                                    add(f"{location}/表格R{ri}C{ci}", cell.text)
                    elif shape.has_text_frame:
                        add(location, shape.text)
            for index, slide in enumerate(presentation.slides, 1):
                shapes(slide.shapes, f"幻灯片{index}")
            warnings.append("按形状顺序提取文字与表格；未包含备注、图表、SmartArt、图片文字或音视频。")
    if not sections:
        raise ExtractionError("未提取到可用文本；扫描件需要 OCR，当前未执行。")
    return {"parserVersion": VERSION, "format": kind, "sections": sections,
            "warnings": warnings,
            "text": "\n\n".join(f"[{s['locator']}]\n{s['text']}" for s in sections)}


def main():
    # Parent must impose a wall-clock timeout and bounded stdout as well.
    import logging
    import resource
    import warnings
    logging.disable(logging.CRITICAL)
    warnings.simplefilter("ignore")
    resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
    if sys.platform.startswith("linux"):
        resource.setrlimit(resource.RLIMIT_AS, (1024**3, 1024**3))
    try:
        if len(sys.argv) != 2:
            raise ExtractionError("缺少文档格式。")
        data = sys.stdin.buffer.read(MAX_BYTES + 1)
        result = extract(data, sys.argv[1])
        print(json.dumps(result, ensure_ascii=False))
    except ExtractionError as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False))
        return 2
    except Exception:
        # Never return parser traces or document bytes to the product error path.
        print(json.dumps({"error": "文件损坏、格式不受支持或无法可靠解析。"}, ensure_ascii=False))
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
