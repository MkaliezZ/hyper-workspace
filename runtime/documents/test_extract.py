import io
import json
import subprocess
import sys
import unittest
import zipfile
from pathlib import Path
from extract import extract, ExtractionError, MAX_BYTES


def saved(obj):
    output = io.BytesIO()
    obj.save(output)
    return output.getvalue()


def pdf(text=True, encrypted=False):
    from pypdf import PdfWriter
    from pypdf.generic import (DictionaryObject, NameObject, DecodedStreamObject)
    writer = PdfWriter()
    page = writer.add_blank_page(300, 300)
    if text:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 100 Td (Synthetic PDF source) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("synthetic-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class ExtractionTests(unittest.TestCase):
    def test_pdf_source_page_and_unreadable_states(self):
        value = extract(pdf(), "pdf")
        self.assertIn("Synthetic PDF source", value["text"])
        self.assertEqual(value["sections"][0]["locator"], "第 1 页")
        for data in [pdf(False), pdf(encrypted=True), b"not a pdf"]:
            with self.assertRaises(ExtractionError):
                extract(data, "pdf")

    def test_docx_preserves_block_and_nested_table_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph("First synthetic paragraph")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "cell-a"
        table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "nested-b"
        document.add_paragraph("Last synthetic paragraph")
        value = extract(saved(document), "docx")
        text = value["text"]
        self.assertLess(text.index("First synthetic"), text.index("cell-a"))
        self.assertLess(text.index("nested-b"), text.index("Last synthetic"))
        self.assertIn("表格R1C2", text)
        self.assertTrue(value["warnings"])

    def test_xlsx_coordinates_formula_and_hidden_sheet(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Source"
        sheet["B2"] = 23
        sheet["C3"] = "=B2+1"
        hidden = workbook.create_sheet("Hidden")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "hidden-sentinel"
        value = extract(saved(workbook), "xlsx")
        self.assertIn("[Source!B2]\n23", value["text"])
        self.assertIn("公式（未计算）：=B2+1", value["text"])
        self.assertNotIn("hidden-sentinel", value["text"])
        self.assertTrue(any("Hidden" in w for w in value["warnings"]))

    def test_pptx_slide_shape_and_table_locations(self):
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "Synthetic slide"
        slide.shapes.add_table(1, 1, 0, Inches(2), Inches(3), Inches(1)).table.cell(0, 0).text = "Table fact"
        value = extract(saved(document), "pptx")
        self.assertIn("Synthetic slide", value["text"])
        self.assertIn("幻灯片1/形状2/表格R1C1", value["text"])
        self.assertIn("Table fact", value["text"])

    def test_format_limits_and_damaged_packages(self):
        for data, kind in [(b"x", "exe"), (b"", "pdf"),
                           (b"x" * (MAX_BYTES + 1), "pdf"), (b"not a zip", "docx")]:
            with self.assertRaises(ExtractionError):
                extract(data, kind)
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active["A2001"] = "outside limit"
        with self.assertRaisesRegex(ExtractionError, "2000"):
            extract(saved(workbook), "xlsx")

    def test_package_guards_do_not_trust_xml_paths_or_compression(self):
        for name, content in [("../word/document.xml", b"x"),
                              ("word/alternate.XML", b"<!DOCTYPE d><d/>"),
                              ("word/custom.data", b"<!DOCTYPE d><d/>"),
                              ("word/document.xml", b"<!DOCTYPE d [<!ENTITY e 'x'>]><d/>"),
                              ("word/document.xml", b"<w:ins xmlns:w='x'>text</w:ins>"),
                              ("word/document.xml", b"x" * (9 * 1024 * 1024)),
                              ("word/vbaProject.bin", b"macro")]:
            output = io.BytesIO()
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("[Content_Types].xml", "<Types/>")
                if name != "word/document.xml":
                    archive.writestr("word/document.xml", "<document/>")
                archive.writestr(name, content)
            with self.assertRaises(ExtractionError):
                extract(output.getvalue(), "docx")

    def test_text_limit_is_explicit_not_silent_truncation(self):
        from docx import Document
        document = Document()
        import random
        text = "".join(random.Random(71).choices("abcdefghijklmnopqrstuvwxyz0123456789", k=20000))
        document.add_paragraph(text)
        result = extract(saved(document), "docx")
        self.assertIn(text, result["text"])
        document.add_paragraph("".join(random.Random(72).choices("abcdefghijklmnopqrstuvwxyz0123456789", k=180001)))
        with self.assertRaisesRegex(ExtractionError, "180000"):
            extract(saved(document), "docx")
        emoji = Document()
        emoji.add_paragraph("".join(random.Random(73).choices(["😀", "🚀", "🌏", "🎉", "🧠", "🤖", "📚", "🔍", "💡", "🛠"], k=90001)))
        with self.assertRaisesRegex(ExtractionError, "180000"):
            extract(saved(emoji), "docx")

    def test_cli_returns_bounded_structured_errors_and_results(self):
        script = Path(__file__).with_name("extract.py")
        for data, success in [(pdf(), True), (b"private-invalid-file-sentinel", False)]:
            result = subprocess.run([sys.executable, str(script), "pdf"], input=data,
                                    capture_output=True, timeout=20)
            self.assertEqual(result.returncode, 0 if success else 2)
            value = json.loads(result.stdout)
            if success:
                self.assertEqual(value["parserVersion"], "document-text-v1")
            else:
                self.assertIn("error", value)
                self.assertNotIn(b"private-invalid", result.stdout + result.stderr)


if __name__ == "__main__":
    unittest.main()
